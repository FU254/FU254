"""Real tiny workbook/DICOM cache tests, with source and provenance changes."""

import copy
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest import mock

import numpy as np
import pandas as pd
import pydicom
from pydicom.dataset import FileDataset, FileMetaDataset
from pydicom.uid import ExplicitVRLittleEndian, SecondaryCaptureImageStorage, generate_uid

from foundation_sft import data
from foundation_sft.cache import load_json, save_json
from foundation_sft.cohort_cache import cached_cohort
from foundation_sft.split import load_or_create_split


class CohortCacheTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="cohort-cache-")
        self.root = Path(self.temporary.name)
        self.dicom_dir = self.root / "PRIVATE_DICOM_DIRECTORY"
        self.dicom_dir.mkdir()
        self.excel = self.root / "PRIVATE_WORKBOOK.xlsx"
        self.cache_dir = self.root / "cache"
        self.config = {
            "excel_path": str(self.excel), "dicom_dir": str(self.dicom_dir),
            "accession_column": "accession", "label_column": "label", "patient_column": "PatientID",
            "patient_id_source": "dicom", "label_definition": "synthetic binary definition",
        }
        self.workbook_rows = [
            {"accession": f"A{index}", "PatientID": f"P{index}", "label": index % 2}
            for index in range(6)
        ]
        self.workbook()
        self.paths = [self.dicom(f"PRIVATE_FILE_{index}.dcm", f"A{index}", f"P{index}") for index in range(6)]

    def tearDown(self):
        self.temporary.cleanup()

    def workbook(self):
        pd.DataFrame(self.workbook_rows).to_excel(self.excel, index=False)

    def dicom(self, filename, accession, patient):
        path = self.dicom_dir / filename
        path.parent.mkdir(parents=True, exist_ok=True)
        meta = FileMetaDataset()
        meta.TransferSyntaxUID = ExplicitVRLittleEndian
        meta.MediaStorageSOPClassUID = SecondaryCaptureImageStorage
        meta.MediaStorageSOPInstanceUID = generate_uid()
        dataset = FileDataset(str(path), {}, file_meta=meta, preamble=b"\0" * 128)
        dataset.SOPClassUID = meta.MediaStorageSOPClassUID
        dataset.SOPInstanceUID = meta.MediaStorageSOPInstanceUID
        dataset.AccessionNumber = accession
        dataset.PatientID = patient
        dataset.Rows = dataset.Columns = 4
        dataset.SamplesPerPixel = 1
        dataset.PhotometricInterpretation = "MONOCHROME2"
        dataset.BitsAllocated = dataset.BitsStored = 16
        dataset.HighBit = 15
        dataset.PixelRepresentation = 0
        dataset.PixelData = np.arange(16, dtype=np.uint16).tobytes()
        self.write(path, dataset)
        return path

    @staticmethod
    def write(path, dataset):
        options = ({"enforce_file_format": True} if int(pydicom.__version__.split(".")[0]) >= 3
                   else {"write_like_original": False})
        pydicom.dcmwrite(path, dataset, **options)

    @staticmethod
    def builder(config):
        # The integration wrapper may land while this independent module is
        # tested. Always exercise the underlying real parser, never recurse.
        return getattr(data, "_build_cohort_uncached", data.build_cohort)(config)

    def cached(self, config=None, builder=None):
        return cached_cohort(self.config if config is None else config, self.cache_dir, self.builder if builder is None else builder)

    def assert_safe(self, message):
        for private in (str(self.root), "PRIVATE_FILE", "PRIVATE_WORKBOOK", "PRIVATE_DICOM_DIRECTORY", "P0", "A0"):
            self.assertNotIn(private, str(message))

    def test_unchanged_hit_skips_workbook_and_dicom_header_parsing(self):
        with mock.patch("foundation_sft.data.pydicom.dcmread", wraps=pydicom.dcmread) as read:
            first = self.cached()
            self.assertEqual(read.call_count, 6)
            self.assertFalse(first.attrs["cohort_cache_hit"])
            with mock.patch("foundation_sft.data.pd.read_excel", side_effect=AssertionError("Unexpected workbook parse")):
                second = self.cached(builder=mock.Mock(side_effect=AssertionError("Unexpected cohort rebuild")))
            self.assertEqual(read.call_count, 6)
        self.assertTrue(second.attrs["cohort_cache_hit"])
        first.attrs.pop("cohort_cache_hit")
        second.attrs.pop("cohort_cache_hit")
        pd.testing.assert_frame_equal(first, second)
        self.assertEqual(first.attrs, second.attrs)

    def test_nested_uppercase_additions_and_removals_invalidate_full_inventory(self):
        first = self.cached()
        directory_mtime = self.dicom_dir.stat().st_mtime_ns
        added = self.dicom("nested/UPPER.DCM", "A0", "P0")
        os.utime(self.dicom_dir, ns=(self.dicom_dir.stat().st_atime_ns, directory_mtime))
        second = self.cached()
        self.assertFalse(second.attrs["cohort_cache_hit"])
        self.assertEqual(len(second), len(first) + 1)
        added.unlink()
        third = self.cached()
        # Exact old source inventories can reuse their immutable old cache.
        self.assertTrue(third.attrs["cohort_cache_hit"])
        self.assertEqual(len(third), len(first))
        self.paths[0].unlink()
        fourth = self.cached()
        self.assertFalse(fourth.attrs["cohort_cache_hit"])
        self.assertEqual(len(fourth), len(first) - 1)

    def test_unmatched_unreadable_files_are_in_inventory(self):
        self.cached()
        extra = self.dicom_dir / "unreadable.dcm"
        extra.write_bytes(b"not a DICOM")
        second = self.cached()
        self.assertFalse(second.attrs["cohort_cache_hit"])
        self.assertEqual(second.attrs["cohort_diagnostics"]["unreadable_dicoms"], 1)
        extra.write_bytes(b"still not a DICOM")
        third = self.cached()
        self.assertFalse(third.attrs["cohort_cache_hit"])

    def test_same_size_and_mtime_header_rewrite_uses_ctime_and_frozen_split_rejects(self):
        first = self.cached()
        manifest = self.root / "frozen.csv"
        load_or_create_split(first, manifest)
        original_stat = self.paths[0].stat()
        dicom = pydicom.dcmread(self.paths[0])
        dicom.PatientID = "Z0"
        self.write(self.paths[0], dicom)
        os.utime(self.paths[0], ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))
        self.assertEqual(self.paths[0].stat().st_size, original_stat.st_size)
        self.assertEqual(self.paths[0].stat().st_mtime_ns, original_stat.st_mtime_ns)
        second = self.cached()
        self.assertFalse(second.attrs["cohort_cache_hit"])
        self.assertIn("Z0", set(second.patient_id))
        with self.assertRaisesRegex(ValueError, "Frozen split manifest"):
            load_or_create_split(second, manifest)

    def test_size_change_rebuilds_and_split_rejects_stale_dicom_stat(self):
        first = self.cached()
        manifest = self.root / "frozen.csv"
        load_or_create_split(first, manifest)
        with self.paths[0].open("ab") as handle:
            handle.write(b"\0\0")
        second = self.cached()
        self.assertFalse(second.attrs["cohort_cache_hit"])
        self.assertNotEqual(set(first.dicom_size), set(second.dicom_size))
        with self.assertRaisesRegex(ValueError, "Frozen split manifest"):
            load_or_create_split(second, manifest)

    def test_workbook_label_change_rebuilds_and_split_rejects(self):
        first = self.cached()
        manifest = self.root / "frozen.csv"
        load_or_create_split(first, manifest)
        self.workbook_rows[0]["label"] = 1
        self.workbook()
        second = self.cached()
        self.assertFalse(second.attrs["cohort_cache_hit"])
        self.assertNotEqual(first.attrs["excel_sha256"], second.attrs["excel_sha256"])
        with self.assertRaisesRegex(ValueError, "Frozen split manifest"):
            load_or_create_split(second, manifest)

    def test_hidden_workbook_content_change_invalidate_even_if_stats_replayed(self):
        from openpyxl import load_workbook
        from foundation_sft.cohort_cache import file_signature

        first = self.cached()
        workbook_stat = file_signature(self.excel)
        workbook = load_workbook(self.excel)
        hidden = workbook.create_sheet("hidden_provenance")
        hidden["A1"] = "new metadata"
        hidden.sheet_state = "hidden"
        workbook.save(self.excel)
        real_signature = file_signature

        def replay_workbook_signature(path):
            return workbook_stat if Path(path) == self.excel else real_signature(path)

        with mock.patch("foundation_sft.cohort_cache.file_signature", side_effect=replay_workbook_signature):
            second = self.cached()
        self.assertFalse(second.attrs["cohort_cache_hit"])
        self.assertNotEqual(first.attrs["excel_sha256"], second.attrs["excel_sha256"])
        pd.testing.assert_frame_equal(first, second)

    def test_configuration_or_runtime_change_is_a_miss(self):
        self.cached()
        config = copy.deepcopy(self.config)
        config["label_definition"] = "different protocol annotation"
        self.assertFalse(self.cached(config).attrs["cohort_cache_hit"])
        with mock.patch("foundation_sft.cohort_cache.runtime_signature", return_value={"changed_runtime": True}):
            self.assertFalse(self.cached().attrs["cohort_cache_hit"])

    def test_build_wrapper_reuses_cohort_across_model_training_settings(self):
        config = {"data": copy.deepcopy(self.config), "cache_dir": str(self.cache_dir),
                  "model": "eva_x", "training": {"epochs": 2}}
        with mock.patch("foundation_sft.data.pydicom.dcmread", wraps=pydicom.dcmread) as read:
            first = data.build_cohort(config)
            self.assertFalse(first.attrs["cohort_cache_hit"])
            self.assertEqual(read.call_count, 6)
            config["model"] = "rad_dino"
            config["training"]["epochs"] = 300
            second = data.build_cohort(config)
            self.assertTrue(second.attrs["cohort_cache_hit"])
            self.assertEqual(read.call_count, 6)
        manifest = self.root / "frozen.csv"
        before = load_or_create_split(first, manifest)
        after = load_or_create_split(second, manifest)
        pd.testing.assert_frame_equal(before, after)

    def test_typed_attrs_columns_dtypes_and_index_survive_reuse_without_aliasing(self):
        def extra_attrs(config):
            frame = self.builder(config)
            frame.attrs["nested"] = {"tuple": ("alpha", 7, True), "list": [{"key": "value"}]}
            frame["nullable"] = pd.Series([pd.NA, 7, 6, 5, 4, 3], dtype="Int64")
            frame.index.name = "samples"
            return frame

        first = self.cached(builder=extra_attrs)
        second = self.cached()
        self.assertTrue(second.attrs["cohort_cache_hit"])
        pd.testing.assert_frame_equal(first, second)
        self.assertEqual(first.attrs["nested"], second.attrs["nested"])
        second.attrs["nested"]["list"][0]["key"] = "changed"
        third = self.cached()
        self.assertEqual(third.attrs["nested"]["list"][0]["key"], "value")
        self.assertEqual(list(first.columns), list(third.columns))
        self.assertEqual(list(first.dtypes), list(third.dtypes))

    def test_corrupt_incompatible_and_missing_payloads_rebuild(self):
        self.cached()
        entries = list((self.cache_dir / "cohorts").glob("*.json"))
        self.assertEqual(len(entries), 1)
        entry = entries[0]
        entry.write_text("{broken JSON", encoding="utf-8")
        self.assertFalse(self.cached().attrs["cohort_cache_hit"])
        envelope = json.loads(entry.read_text())
        envelope["payload"]["cohort"]["dtypes"][0] = "not-a-pandas-dtype"
        entry.write_text(json.dumps(envelope), encoding="utf-8")
        self.assertFalse(self.cached().attrs["cohort_cache_hit"])
        payload = load_json(entry)
        payload["cohort"]["dtypes"][0] = "not-a-pandas-dtype"
        save_json(entry, payload)
        self.assertFalse(self.cached().attrs["cohort_cache_hit"])
        payload = load_json(entry)
        payload["version"] = "future-incompatible-version"
        save_json(entry, payload)
        self.assertFalse(self.cached().attrs["cohort_cache_hit"])
        entry.unlink()
        self.assertFalse(self.cached().attrs["cohort_cache_hit"])

    def test_source_change_during_build_is_refused_without_saved_entry(self):
        def change_during_build(config):
            frame = self.builder(config)
            with self.paths[0].open("ab") as handle:
                handle.write(b"\0\0")
            return frame

        with self.assertRaisesRegex(ValueError, "Cohort sources changed") as error:
            self.cached(builder=change_during_build)
        self.assert_safe(error.exception)
        self.assertEqual(list((self.cache_dir / "cohorts").glob("*.json")), [])

    def test_source_change_during_hit_is_refused(self):
        import foundation_sft.cohort_cache as module
        self.cached()
        real_restore = module._restore_cohort

        def change_on_restore(payload):
            frame = real_restore(payload)
            with self.paths[0].open("ab") as handle:
                handle.write(b"\0\0")
            return frame

        with mock.patch.object(module, "_restore_cohort", side_effect=change_on_restore):
            with self.assertRaisesRegex(ValueError, "Cohort sources changed") as error:
                self.cached()
        self.assert_safe(error.exception)

    def test_missing_sources_are_sanitized_and_cache_files_private(self):
        self.cached()
        directory = self.cache_dir / "cohorts"
        self.assertEqual(directory.stat().st_mode & 0o777, 0o700)
        entries = list(directory.glob("*.json"))
        self.assertEqual(entries[0].stat().st_mode & 0o777, 0o600)
        self.assertRegex(entries[0].name, r"^[0-9a-f]{64}\.json$")
        self.excel.unlink()
        with self.assertRaises(FileNotFoundError) as error:
            self.cached()
        self.assert_safe(error.exception)


if __name__ == "__main__":
    unittest.main()
